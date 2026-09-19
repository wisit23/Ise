from openpyxl import load_workbook

path = r"D:\ISE Project\Ise\outputs\test-cases4-corrected-20260915\Test_Cases4_Final_corrected.xlsx"
workbook = load_workbook(path)
removed = 0
for sheet in workbook.worksheets:
    for row in sheet.iter_rows():
        for cell in row:
            if cell.comment is not None:
                cell.comment = None
                removed += 1
workbook.save(path)
print(f"removed_comments={removed}")
