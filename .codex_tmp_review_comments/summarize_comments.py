import json
import re
import sys
from collections import Counter, defaultdict

from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")
path = r"C:\Users\sirad\Downloads\Test_Cases4_Final_reviewed.xlsx"
wb = load_workbook(path, read_only=False, data_only=False)

items = []
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for cell in row:
            if cell.comment is None:
                continue
            text = cell.comment.text or ""
            if text.startswith("🔴"):
                severity = "red"
            elif text.startswith("🟠"):
                severity = "orange"
            elif text.startswith("🟡"):
                severity = "yellow"
            else:
                severity = "unmarked"
            first_line = text.splitlines()[0].strip()
            items.append({
                "sheet": ws.title,
                "cell": cell.coordinate,
                "row": cell.row,
                "test_case_id": ws.cell(cell.row, 1).value,
                "test_case": ws.cell(cell.row, 3).value,
                "severity": severity,
                "first_line": first_line,
                "text": text,
                "author": cell.comment.author,
            })

by_severity = Counter(i["severity"] for i in items)
cells_by_severity = defaultdict(list)
for i in items:
    cells_by_severity[i["severity"]].append(i["cell"])

print(json.dumps({
    "total": len(items),
    "by_severity": by_severity,
    "cells_by_severity": cells_by_severity,
    "items": items,
}, ensure_ascii=False, indent=2))
